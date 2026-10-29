"""Import an explicitly downloaded K-FIND XLSX. Requires openpyxl only at import time.

The source workbook stays outside Git. Volume-only values are not silently treated as grams.
"""

import argparse
import collections
import csv
import hashlib
import json
import math
from pathlib import Path

FIELDS = {
    "energy_kcal": "에너지(kcal)",
    "carb_g": "탄수화물(g)",
    "sugar_g": "당류(g)",
    "protein_g": "단백질(g)",
    "fat_g": "지방(g)",
    "sat_fat_g": "포화지방산(g)",
    "sodium_mg": "나트륨(mg)",
    "fiber_g": "식이섬유(g)",
    "potassium_mg": "칼륨(mg)",
    "cholesterol_mg": "콜레스테롤(mg)",
}


def number(value):
    if value is None or not str(value).strip():
        return None
    try:
        result = float(str(value).replace(",", ""))
    except ValueError:
        return None
    return result if math.isfinite(result) and result >= 0 else None


def main():
    import openpyxl

    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--published-on", required=True)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[3]
    sha = hashlib.sha256(args.input.read_bytes()).hexdigest()
    version = f"mfds-kddb-{args.published_on.replace('-', '')}-{sha[:8]}"
    sheet = openpyxl.load_workbook(args.input, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows))
    required = {"식품코드", "식품명", "영양성분함량기준량", *FIELDS.values()}
    if not required <= set(headers):
        raise ValueError("Unexpected MFDS schema; do not guess column positions")
    foods = []
    skipped = collections.Counter()
    missing = collections.Counter()
    seen = set()
    for values in rows:
        row = dict(zip(headers, values))
        if str(row["영양성분함량기준량"]).lower().replace(" ", "") != "100g":
            skipped[str(row["영양성분함량기준량"])] += 1
            continue
        code = str(row["식품코드"])
        if code in seen:
            raise ValueError("Duplicate source food code")
        seen.add(code)
        serving = str(row.get("식품중량") or "").strip().lower()
        serving_g = number(serving[:-1]) if serving.endswith("g") else None
        nutrients = {key: number(row.get(column)) for key, column in FIELDS.items()}
        missing.update(key for key, value in nutrients.items() if value is None)
        foods.append(
            {
                "food_code": code,
                "name": str(row["식품명"]),
                "category": str(row.get("식품대분류명") or ""),
                "serving_g": serving_g,
                "origin": str(row.get("식품기원명") or ""),
                "source_date": str(row.get("데이터기준일자") or ""),
                "per_100g": nutrients,
            }
        )
    if not foods:
        raise ValueError("No gram-based foods")
    # Deterministic priority: measured household / measured restaurant, then other sources.
    foods.sort(
        key=lambda food: (
            not ("분석" in food["origin"] and "재료량" not in food["origin"]),
            food["food_code"],
        )
    )
    meta = {
        "version": version,
        "kind": "official_mfds",
        "publisher": "식품의약품안전처",
        "title": "K-FIND 음식 DB",
        "source_url": "https://various.foodsafetykorea.go.kr/nutrient/general/down/historyList.do",
        "published_on": args.published_on,
        "downloaded_on": "2026-10-29",
        "source_sha256": sha,
        "source_rows": sheet.max_row - 1,
        "food_count": len(foods),
        "excluded_basis": dict(skipped),
        "missing_nutrients": dict(missing),
        "basis": "100g",
        "note": "100ml 자료는 밀도 정보 없이 g으로 바꾸지 않고 제외. 결측·검출한계 표시는 0으로 채우지 않음.",
    }
    target = root / "services/diet-f6/data"
    with (target / "food_nutrition_mfds.csv").open(
        "w", encoding="utf-8", newline=""
    ) as file:
        keys = [
            "food_code",
            "name",
            "category",
            "serving_g",
            "origin",
            "source_date",
            *FIELDS,
        ]
        writer = csv.DictWriter(file, fieldnames=keys)
        writer.writeheader()
        for food in foods:
            writer.writerow(
                {
                    **{k: v for k, v in food.items() if k != "per_100g"},
                    **food["per_100g"],
                }
            )
    (target / "catalog-source.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    frontend = root / "frontend/src/design"
    (frontend / "mfds-foods.json").write_text(
        json.dumps(foods, ensure_ascii=False, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )
    (frontend / "food-catalog-source.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(
        json.dumps(
            {"version": version, "foods": len(foods), "excludedBasis": dict(skipped)},
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
