"""
[파일] analysis_pack.py
[용도] 데이터 분석 묶음. 최신 explore_*(단계별 히스토그램·Spearman 상관 행렬·구성비) 와 feature_importance_*(변수별 모델 기여도) 결과를
      엑셀 한 권(data_analysis.xlsx, 시트 8개)으로 모으고 그림 5장을 같은 폴더(reports/analysis_<시각>/)에 복사한다. 계산은 하지 않는다.
[왜 필요한가] 멘토링 이후 데이터 분석을 다시 시작하면서(2026-10-23) 변수별 영향 크기·분포·상관을 한 파일로 보기 위해.
[쓰는 곳] 사용자가 직접 실행. 먼저 feature_importance.py 와 explore_data.py 를 돌려 둔다.
      python3 analysis_pack.py [--explore reports/explore_...] [--importance reports/feature_importance_...]
"""

from __future__ import annotations

import argparse
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import thresholds as T

MODELS = ["ordinal__base", "nested__base", "lgbm__base"]
SHORT = {"ordinal__base": "ordinal", "nested__base": "nested", "lgbm__base": "lgbm"}
HEAD_FILL = PatternFill("solid", fgColor="E8EEF7")
PNGS = {"hist_by_stage.png": "explore", "correlation.png": "explore", "stage_share_by_category.png": "explore", "stage_share_by_group.png": "explore", "feature_importance.png": "importance"}


def latest(prefix: str, base: Path) -> Path:
    dirs = sorted(d for d in base.glob(f"{prefix}_*") if d.is_dir())
    if not dirs:
        raise SystemExit(f"{base}/{prefix}_* 폴더가 없다. 먼저 해당 스크립트를 실행한다.")
    return dirs[-1]


def wide(dir_: Path, kind: str, cols: list[str]) -> pd.DataFrame:
    """features_/groups_ csv 세 모델을 열로 나란히(share_ordinal, share_nested, ...). 평균 share 순."""
    parts = {}
    for m in MODELS:
        f = dir_ / f"{kind}_{m}.csv"
        if f.exists():
            parts[m] = pd.read_csv(f, index_col=0)
    out = pd.DataFrame(index=sorted({i for t in parts.values() for i in t.index}))
    for c in cols:
        for m, t in parts.items():
            if c in t.columns:
                out[f"{c}_{SHORT[m]}"] = t[c]
    extra = {"ordinal__base": "coef", "nested__base": "coef_mean3", "lgbm__base": "gain_share"}
    for m, c in extra.items():
        if m in parts and c in parts[m].columns:
            out[f"{c}_{SHORT[m]}"] = parts[m][c]
    share_cols = [c for c in out.columns if c.startswith("share_")]
    out["share_mean"] = out[share_cols].mean(axis=1)
    out = out.sort_values("share_mean", ascending=False)
    out.index.name = "feature" if kind == "features" else "code"
    return out.reset_index()


def style(ws, nd: int = 3, freeze: str = "B2") -> None:
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = freeze
    for col in ws.iter_cols(min_row=2):
        for c in col:
            if isinstance(c.value, float):
                c.number_format = "0." + "0" * nd
    for i, col in enumerate(ws.columns, 1):
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(9, w * 1.1 + 2), 44)


def main(argv=None) -> Path:
    ap = argparse.ArgumentParser(description="탐색·기여도 결과를 엑셀 한 권으로 묶는다")
    ap.add_argument("--explore", default=None)
    ap.add_argument("--importance", default=None)
    ap.add_argument("--out", default="reports")
    a = ap.parse_args(argv)
    base = Path(a.out)
    ex = Path(a.explore) if a.explore else latest("explore", base)
    im = Path(a.importance) if a.importance else latest("feature_importance", base)
    out = base / f"analysis_{datetime.now():%Y%m%d_%H%M%S}"
    out.mkdir(parents=True, exist_ok=True)

    feat = wide(im, "features", ["share", "mean_abs", "mean"])
    grp = wide(im, "groups", ["share", "mean_abs", "mean"])
    corr = pd.read_csv(ex / "correlation_matrix.csv", index_col=0)
    stage_corr = corr["stage_index"].drop("stage_index")
    stage_corr = stage_corr.reindex(stage_corr.abs().sort_values(ascending=False).index).rename("spearman_with_stage").reset_index().rename(columns={"index": "variable"})
    med = pd.read_csv(ex / "stage_medians.csv")
    cat = pd.read_csv(ex / "category_shares.csv")
    grp_share = pd.read_csv(ex / "group_shares.csv")

    readme = pd.DataFrame({"시트": ["기여도_변수", "기여도_이유코드", "상관행렬", "단계상관", "단계별_중앙값", "범주별_구성비", "집단별_구성비"],
                           "내용": ["모델 3종의 변수별 전체 기여도. 내부 테스트 분할에서 행별 로짓 기여도(선형: 계수×표준화값, LightGBM: SHAP; nested·lgbm 은 누적 이진 3개 합)를 모아 평균. share = 전체 절대 기여도 중 비중, mean_abs = 평균 절대 기여도(영향 크기), mean = 평균 기여도(양수면 평균적으로 단계 확률을 올림). coef/coef_mean3 = 표준화 계수, gain_share = LightGBM gain 비중. share_mean 순 정렬.",
                                    "같은 기여도를 이유 코드 단위(코드가 보는 피처 묶음)로 합한 것. sex(male) 은 코드가 아니지만 비교용으로 포함.",
                                    "학습 데이터(범위·결측 검증 통과 행)의 Spearman 상관 행렬. sbp/dbp (label only) 는 라벨을 만든 측정 혈압이며 피처가 아님. stage_index 는 단계 0~3. 상관은 단조 관계의 세기만 보여주고 인과를 뜻하지 않음.",
                                    "stage_index 와의 Spearman 상관을 절대값 순으로 정렬.",
                                    "단계별 연속 변수 중앙값(히스토그램 hist_by_stage.png 의 수치 요약).",
                                    "범주형 변수 수준별 단계 구성비(행 비율)와 행 수.",
                                    "성별×연령대, BMI 구간별 단계 구성비."],
                           "출처": [str(im), str(im), str(ex), str(ex), str(ex), str(ex), str(ex)]})
    sheets = {"설명": readme, "기여도_변수": feat, "기여도_이유코드": grp, "상관행렬": corr.reset_index().rename(columns={"index": "variable"}),
              "단계상관": stage_corr, "단계별_중앙값": med, "범주별_구성비": cat, "집단별_구성비": grp_share}
    xlsx = out / "data_analysis.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
        for name, df in sheets.items():
            df.to_excel(w, sheet_name=name, index=False)
        wb = w.book
        for name in sheets:
            style(wb[name], nd=1 if name == "단계별_중앙값" else 3, freeze="A2" if name == "설명" else "B2")
        ws = wb["설명"]
        ws.column_dimensions["B"].width = 110
        for row in ws.iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
        ws = wb["상관행렬"]
        n = len(corr)
        rng = f"B2:{get_column_letter(n + 1)}{n + 1}"
        ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=-1, start_color="E34948", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="num", end_value=1, end_color="2A78D6"))
        for name, df in (("기여도_변수", feat), ("기여도_이유코드", grp)):
            ws = wb[name]
            for j, c in enumerate(df.columns, 1):
                if c.startswith("share_"):
                    ws.conditional_formatting.add(f"{get_column_letter(j)}2:{get_column_letter(j)}{len(df) + 1}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=0.5, color="5598E7"))
    for png, src in PNGS.items():
        f = (ex if src == "explore" else im) / png
        if f.exists():
            shutil.copy(f, out / png)
    print(f"완료: {xlsx} (시트 {len(sheets)}개), 그림 {sum((out / p).exists() for p in PNGS)}장  ← {ex.name}, {im.name}")
    return out


if __name__ == "__main__":
    main()
